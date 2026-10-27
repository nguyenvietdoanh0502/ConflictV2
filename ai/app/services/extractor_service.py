import io
import logging
from app.errors import AIEngineError
from app.schemas import ExtractTextResponse

logger = logging.getLogger("unisync-ai")


def extract_text_from_file(file_bytes: bytes, filename: str) -> ExtractTextResponse:
    """Trích xuất văn bản từ file PDF, DOCX hoặc TXT.
    
    Phát hiện và báo lỗi rõ ràng nếu file rỗng hoặc file PDF là bản scan ảnh
    không có lớp chữ (tránh để AI hallucinate nội dung theo AGENT.md §8.3).
    """
    if not file_bytes or len(file_bytes) == 0:
        raise AIEngineError(
            error_code="EMPTY_FILE",
            message="File tải lên rỗng (0 bytes). Vui lòng chọn tài liệu hợp lệ.",
            status_code=400
        )

    fn_lower = filename.lower()
    extracted_text = ""
    page_count = 1

    try:
        if fn_lower.endswith(".pdf"):
            import pymupdf
            doc = pymupdf.open(stream=file_bytes, filetype="pdf")
            page_count = len(doc)
            pages_text = []
            for page in doc:
                text = page.get_text("text")
                if text:
                    pages_text.append(text.strip())
            extracted_text = "\n\n".join(pages_text)
            doc.close()

            # Kiểm tra nếu PDF scan ảnh không có text
            if len(extracted_text.strip()) < 50:
                raise AIEngineError(
                    error_code="MATERIAL_PARSE_FAILED",
                    message="Tài liệu PDF không chứa văn bản hoặc là tài liệu scan dạng hình ảnh. Vui lòng sử dụng file PDF có văn bản kỹ thuật số.",
                    status_code=422
                )

        elif fn_lower.endswith(".docx"):
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            parts = []
            for p in doc.paragraphs:
                if p.text and p.text.strip():
                    parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        parts.append(" | ".join(row_cells))
            extracted_text = "\n".join(parts)
            page_count = max(1, len(parts) // 25)

        elif fn_lower.endswith(".txt"):
            for encoding in ("utf-8", "utf-8-sig", "cp1258", "latin-1"):
                try:
                    extracted_text = file_bytes.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            page_count = 1

        else:
            raise AIEngineError(
                error_code="INVALID_FILE_TYPE",
                message=f"Định dạng file không được hỗ trợ: '{filename}'. Hệ thống chỉ chấp nhận file .pdf, .docx, hoặc .txt.",
                status_code=400
            )

    except AIEngineError:
        raise
    except Exception as e:
        logger.error(f"Lỗi khi trích xuất văn bản từ '{filename}': {e}", exc_info=True)
        raise AIEngineError(
            error_code="MATERIAL_PARSE_FAILED",
            message=f"Không thể đọc nội dung file '{filename}': {str(e)}",
            status_code=422
        )

    clean_text = extracted_text.strip()
    if len(clean_text) < 30:
        raise AIEngineError(
            error_code="MATERIAL_PARSE_FAILED",
            message="Tài liệu không có đủ nội dung văn bản (dưới 30 ký tự) để tạo học liệu ôn tập.",
            status_code=422
        )

    # Giới hạn an toàn khoảng 50,000 ký tự để tập trung vào phần cốt lõi của đề cương
    if len(clean_text) > 50000:
        logger.info(f"Tài liệu '{filename}' dài {len(clean_text)} ký tự, cắt ngắn còn 50,000 ký tự đầu tiên.")
        clean_text = clean_text[:50000]

    return ExtractTextResponse(
        text=clean_text,
        char_count=len(clean_text),
        page_count=page_count
    )
